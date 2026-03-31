"""
╔══════════════════════════════════════════════════════════════════╗
║   AI SmartSend — Versión de PRUEBA (Trial)                      ║
║   Envío automatizado de WhatsApp                                ║
╚══════════════════════════════════════════════════════════════════╝

INSTALACIÓN:
    pip install selenium webdriver-manager schedule openpyxl

EJECUTAR:
    python whatsapp_trial.py

LÍMITE: 10 envíos de prueba gratuitos.
Para la versión completa sin límites, contacta al proveedor.
"""

import sys
import time
import threading
import logging
import tkinter as tk
from tkinter import ttk, messagebox, scrolledtext
from datetime import datetime
from pathlib import Path
from urllib.parse import quote

import openpyxl
import schedule

# ─── Logging ────────────────────────────────────────────────────
logging.basicConfig(
    filename="ai_smartsend.log",
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    encoding="utf-8"
)
log = logging.getLogger("AISmartSend")

CARPETA   = Path(__file__).parent
EXCEL_CON = CARPETA / "contactos.xlsx"
EXCEL_REP = CARPETA / "reporte_envios.xlsx"
SESION    = str(CARPETA / "sesion_chrome")

# ── Módulo de seguridad ──────────────────────────────────────────
_sec_path = str(CARPETA)
if _sec_path not in sys.path:
    sys.path.insert(0, _sec_path)
from security import (
    obtener_o_crear_device_id,
    generar_password_instalacion,
    trial_agotado,
    marcar_trial_agotado,
    leer_contador,
    incrementar_contador,
)

# ── Configuración Trial ──────────────────────────────────────────
MODO_PRUEBA   = True
LIMITE_PRUEBA = 10


def _leer_contador() -> int:
    return leer_contador()

def _incrementar_contador(n: int):
    incrementar_contador(n)

def _envios_restantes() -> int:
    return max(0, LIMITE_PRUEBA - _leer_contador())


# ══════════════════════════════════════════════════════════════════
#  HELPERS
# ══════════════════════════════════════════════════════════════════
def cargar_contactos() -> dict:
    if not EXCEL_CON.exists():
        _crear_excel_ejemplo()
    contactos = {}
    try:
        wb = openpyxl.load_workbook(EXCEL_CON)
        ws = wb.active
        enc = [str(c.value).strip().lower() if c.value else "" for c in ws[1]]
        if "telefono" not in enc:
            return {}
        it = enc.index("telefono")
        ia = enc.index("activo") if "activo" in enc else None
        for fila in ws.iter_rows(min_row=2, values_only=True):
            if not fila[it]:
                continue
            if ia is not None:
                val_activo = str(fila[ia]).strip().upper() if fila[ia] is not None else ""
                if val_activo in ("NO", "FALSE", "0", "N", "NO ", "N ") or val_activo.startswith("N"):
                    continue
            tel = str(fila[it]).strip()
            if not tel.startswith("+"):
                tel = "+" + tel
            datos = {}
            for i, e in enumerate(enc):
                if e and e not in ("telefono", "activo") and i < len(fila):
                    datos[e] = str(fila[i]).strip() if fila[i] is not None else ""
            contactos[tel] = datos
    except Exception as ex:
        log.error(f"Error leyendo Excel: {ex}")
    return contactos


def grupos_disponibles(contactos: dict) -> list:
    gs = sorted(set(d.get("grupo", "").lower() for d in contactos.values() if d.get("grupo")))
    return ["todos"] + gs


def personalizar(texto: str, datos: dict) -> str:
    ahora = datetime.now()
    dias  = ["Lunes", "Martes", "Miércoles", "Jueves", "Viernes", "Sábado", "Domingo"]
    meses = ["Enero", "Febrero", "Marzo", "Abril", "Mayo", "Junio",
             "Julio", "Agosto", "Septiembre", "Octubre", "Noviembre", "Diciembre"]
    texto = texto.replace("{fecha}", ahora.strftime("%d/%m/%Y"))
    texto = texto.replace("{hora}",  ahora.strftime("%H:%M"))
    texto = texto.replace("{dia}",   dias[ahora.weekday()])
    texto = texto.replace("{mes}",   meses[ahora.month - 1])
    for k, v in datos.items():
        texto = texto.replace(f"{{{k}}}", str(v))
    return texto


def matar_chrome():
    """Cierra procesos Chrome/Chromedriver residuales. Multiplataforma."""
    import platform
    import subprocess
    try:
        if platform.system() == "Windows":
            subprocess.run(["taskkill", "/F", "/IM", "chrome.exe"],
                           capture_output=True, timeout=5)
            subprocess.run(["taskkill", "/F", "/IM", "chromedriver.exe"],
                           capture_output=True, timeout=5)
        else:
            subprocess.run(["pkill", "-f", "chrome"],
                           capture_output=True, timeout=5)
            subprocess.run(["pkill", "-f", "chromedriver"],
                           capture_output=True, timeout=5)
        time.sleep(2)
    except Exception:
        pass


def crear_driver():
    from selenium import webdriver
    from selenium.webdriver.chrome.options import Options
    from selenium.webdriver.chrome.service import Service
    from webdriver_manager.chrome import ChromeDriverManager

    ops = Options()
    ops.add_argument(f"--user-data-dir={SESION}")
    ops.add_argument("--profile-directory=Default")
    ops.add_argument("--no-sandbox")
    ops.add_argument("--disable-dev-shm-usage")
    ops.add_argument("--disable-extensions")
    ops.add_argument("--disable-gpu")
    ops.add_argument("--remote-debugging-port=9222")
    ops.add_argument("--no-first-run")
    ops.add_argument("--no-default-browser-check")
    ops.add_experimental_option("excludeSwitches", ["enable-logging"])
    ops.add_experimental_option("useAutomationExtension", False)

    svc = Service(ChromeDriverManager().install())
    driver = webdriver.Chrome(service=svc, options=ops)
    driver.execute_cdp_cmd("Page.addScriptToEvaluateOnNewDocument", {
        "source": "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
    })
    return driver


def _encontrar_boton_enviar(driver, wait):
    """Selector robusto con 6 estrategias de fallback."""
    from selenium.webdriver.common.by import By
    from selenium.webdriver.support import expected_conditions as EC

    estrategias = [
        (By.XPATH, '//button[@aria-label="Enviar"]'),
        (By.XPATH, '//button[@aria-label="Send"]'),
        (By.XPATH, '//button[@data-tab="11"]'),
        (By.XPATH, '//button[contains(@class,"tvf2evcx") or contains(@class,"odfin822")]'),
        (By.XPATH, '//button[.//*[@data-icon="send"]]'),
        (By.XPATH, '//footer//button[last()]'),
    ]
    for by, selector in estrategias:
        try:
            boton = wait.until(EC.element_to_be_clickable((by, selector)))
            log.info(f"  ▸ Botón Enviar encontrado: {selector[:60]}")
            return boton
        except Exception:
            continue
    raise RuntimeError("No se encontró el botón Enviar con ninguna estrategia conocida.")


def enviar_a_contactos(contactos: dict, texto_template: str,
                        callback_log=None, espera: int = 50) -> list:
    from selenium.webdriver.support.ui import WebDriverWait

    resultados = []
    driver = None

    def _log(msg):
        log.info(msg)
        if callback_log:
            callback_log(msg)

    try:
        matar_chrome()
        driver = crear_driver()

        _log("🌐 Cargando WhatsApp Web...")
        driver.get("https://web.whatsapp.com")
        time.sleep(20)

        for tel, datos in contactos.items():
            nombre = datos.get("nombre", tel)
            try:
                _log(f"📤 Enviando a {nombre} ({tel})...")
                mensaje = personalizar(texto_template, datos)
                num = tel.replace("+", "").replace(" ", "")
                url = f"https://web.whatsapp.com/send?phone={num}&text={quote(mensaje)}"
                driver.get(url)
                time.sleep(5)

                wait = WebDriverWait(driver, espera)
                boton = _encontrar_boton_enviar(driver, wait)
                time.sleep(2)
                boton.click()
                time.sleep(4)

                _log(f"✅ Enviado a {nombre}")
                resultados.append({
                    "telefono": tel, "nombre": nombre,
                    "estado": "Enviado",
                    "hora": datetime.now().strftime("%H:%M:%S"),
                    "mensaje": mensaje
                })
            except Exception as e:
                error_msg = str(e)[:120] if str(e) else "Timeout — WhatsApp Web no respondió"
                _log(f"❌ Error con {nombre} ({tel}): {error_msg}")
                resultados.append({
                    "telefono": tel, "nombre": nombre,
                    "estado": f"Fallido: {error_msg}",
                    "hora": datetime.now().strftime("%H:%M:%S"),
                    "mensaje": ""
                })
    except Exception as e:
        _log(f"❌ Error iniciando Chrome: {e}")
    finally:
        if driver:
            driver.quit()
    return resultados


def guardar_reporte(resultados: list):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte"
    enc = ["Fecha", "Hora", "Teléfono", "Nombre", "Estado", "Mensaje"]
    for col, e in enumerate(enc, 1):
        c = ws.cell(row=1, column=col, value=e)
        try:
            from openpyxl.styles import Font, PatternFill
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="075E54")
        except Exception:
            pass
    fecha_hoy = datetime.now().strftime("%d/%m/%Y")
    for row, r in enumerate(resultados, 2):
        ws.cell(row=row, column=1, value=fecha_hoy)
        ws.cell(row=row, column=2, value=r["hora"])
        ws.cell(row=row, column=3, value=r["telefono"])
        ws.cell(row=row, column=4, value=r["nombre"])
        ws.cell(row=row, column=5, value=r["estado"])
        ws.cell(row=row, column=6, value=r["mensaje"])
    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 4, 60)
    wb.save(EXCEL_REP)
    log.info(f"Reporte guardado en {EXCEL_REP}")


def _crear_excel_ejemplo():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Contactos"
    for col, v in enumerate(["telefono", "nombre", "grupo", "activo"], 1):
        ws.cell(row=1, column=col, value=v)
    for row, d in enumerate([
        ["+51987654321", "Juan Pérez",   "clientes", "SI"],
        ["+51912345678", "María García", "clientes", "SI"],
        ["+51999111222", "Carlos López", "equipo",   "SI"],
    ], 2):
        for col, v in enumerate(d, 1):
            ws.cell(row=row, column=col, value=v)
    wb.save(EXCEL_CON)


# ══════════════════════════════════════════════════════════════════
#  INTERFAZ GRÁFICA
# ══════════════════════════════════════════════════════════════════
class App(tk.Tk):

    VERDE_WA  = "#25D366"
    VERDE_OSC = "#075E54"
    VERDE_MED = "#128C7E"
    GRIS_BG   = "#f5f5f5"

    def __init__(self):
        super().__init__()
        self.title("AI SmartSend — Versión de Prueba")
        self.geometry("740x900")
        self.minsize(640, 600)
        self.resizable(True, True)
        self.configure(bg=self.GRIS_BG)
        self.protocol("WM_DELETE_WINDOW", self._salir)

        self.contactos   = cargar_contactos()
        self.grupos_vars = {}
        self._job_activo = False
        self._hilo       = None

        self._construir_ui()
        self._actualizar_grupos()
        self._actualizar_contador()

        # Verificar si el trial ya está agotado al iniciar
        self.after(500, self._verificar_trial)

    def _verificar_trial(self):
        if trial_agotado():
            messagebox.showerror("AI SmartSend — Versión de Prueba",
                "Este equipo ya utilizó su versión de prueba.\n\n"
                "Para continuar usando AI SmartSend, adquiere\n"
                "la versión completa contactando al proveedor.")
            self.destroy()

    # ── UI principal ──────────────────────────────────────────────
    def _construir_ui(self):
        # ── Barra superior ──
        barra = tk.Frame(self, bg=self.VERDE_OSC)
        barra.pack(fill="x")
        lbl_titulo = tk.Label(barra,
                              text="📱  AI SmartSend — Versión de Prueba",
                              font=("Segoe UI", 15, "bold"),
                              bg=self.VERDE_OSC, fg="white", cursor="fleur")
        lbl_titulo.pack(side="left", padx=16, pady=10)
        lbl_titulo.bind("<ButtonPress-1>", self._inicio_arrastre)
        lbl_titulo.bind("<B1-Motion>",     self._arrastrar)
        barra.bind("<ButtonPress-1>",      self._inicio_arrastre)
        barra.bind("<B1-Motion>",          self._arrastrar)
        tk.Button(barra, text="✕  Salir",
                  command=self._salir,
                  bg="#c0392b", fg="white",
                  font=("Segoe UI", 9, "bold"),
                  relief="flat", cursor="hand2",
                  padx=10).pack(side="right", padx=10, pady=8)

        # ── Banner Trial ──
        self.lbl_prueba = tk.Label(
            self,
            text=f"⚠️  VERSIÓN DE PRUEBA — {_envios_restantes()} envíos restantes de {LIMITE_PRUEBA}",
            font=("Segoe UI", 9, "bold"),
            bg="#ffeeba", fg="#856404", anchor="center")
        self.lbl_prueba.pack(fill="x")

        # ── Área con scroll ──
        canvas = tk.Canvas(self, bg=self.GRIS_BG, highlightthickness=0)
        scroll = ttk.Scrollbar(self, orient="vertical", command=canvas.yview)
        self.frame = tk.Frame(canvas, bg=self.GRIS_BG)
        self.frame.bind("<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.create_window((0, 0), window=self.frame, anchor="nw")
        canvas.configure(yscrollcommand=scroll.set)
        canvas.pack(side="left", fill="both", expand=True)
        scroll.pack(side="right", fill="y")
        canvas.bind_all("<MouseWheel>",
            lambda e: canvas.yview_scroll(int(-1*(e.delta/120)), "units"))
        canvas.bind("<ButtonPress-1>", self._inicio_arrastre)
        canvas.bind("<B1-Motion>",     self._arrastrar)

        f = self.frame

        # ══ 1. Mensaje ════════════════════════════════════════════
        self._seccion(f, "1. ✏️   Mensaje a enviar")
        tk.Label(f, text="Variables: {nombre}  {fecha}  {hora}  {dia}  {mes}",
                 font=("Segoe UI", 8), fg="#666", bg=self.GRIS_BG
                 ).pack(anchor="w", padx=14)
        self.txt_mensaje = scrolledtext.ScrolledText(
            f, height=5, width=80, font=("Segoe UI", 10), wrap="word")
        self.txt_mensaje.pack(padx=14, pady=4)
        self.txt_mensaje.insert("1.0",
            "🎉 ¡Hola {nombre}! Hoy es {fecha}. Te escribimos para informarte de nuestras novedades. 🚀")

        # ══ 2. Grupos ═════════════════════════════════════════════
        self._seccion(f, "2. 👥  Grupos destinatarios")
        self.frame_grupos = tk.Frame(f, bg=self.GRIS_BG)
        self.frame_grupos.pack(anchor="w", padx=14, pady=4)
        tk.Button(f, text="🔄  Recargar contactos del Excel",
                  command=self._recargar_contactos,
                  bg=self.VERDE_MED, fg="white",
                  font=("Segoe UI", 9), relief="flat",
                  cursor="hand2", padx=8, pady=4
                  ).pack(anchor="w", padx=14, pady=2)

        # ══ 3. Fecha y hora ═══════════════════════════════════════
        self._seccion(f, "3. 📅  Fecha y hora de envío")
        frame_dt = tk.Frame(f, bg=self.GRIS_BG)
        frame_dt.pack(anchor="w", padx=14, pady=4)

        self.var_inmediato = tk.BooleanVar(value=True)
        tk.Checkbutton(frame_dt, text="⚡ Envío inmediato (ahora mismo)",
                       variable=self.var_inmediato,
                       command=self._toggle_datetime,
                       bg=self.GRIS_BG, font=("Segoe UI", 10)
                       ).grid(row=0, column=0, columnspan=4, sticky="w", pady=4)

        tk.Label(frame_dt, text="Fecha (DD/MM/AAAA):",
                 bg=self.GRIS_BG, font=("Segoe UI", 10)
                 ).grid(row=1, column=0, sticky="w", padx=4, pady=4)
        self.ent_fecha = tk.Entry(frame_dt, width=14, font=("Segoe UI", 10))
        self.ent_fecha.insert(0, datetime.now().strftime("%d/%m/%Y"))
        self.ent_fecha.grid(row=1, column=1, padx=4)

        tk.Label(frame_dt, text="Hora (HH:MM):",
                 bg=self.GRIS_BG, font=("Segoe UI", 10)
                 ).grid(row=1, column=2, sticky="w", padx=4)
        self.ent_hora = tk.Entry(frame_dt, width=8, font=("Segoe UI", 10))
        self.ent_hora.insert(0, datetime.now().strftime("%H:%M"))
        self.ent_hora.grid(row=1, column=3, padx=4)
        self._toggle_datetime()

        # ══ 4. Repetición ═════════════════════════════════════════
        self._seccion(f, "4. 🔁  Repetición del envío")
        frame_rep = tk.Frame(f, bg=self.GRIS_BG)
        frame_rep.pack(anchor="w", padx=14, pady=4)
        self.var_repetir = tk.StringVar(value="diario")
        opciones = [
            ("Diario",              "diario",       0, 0),
            ("Semanal",             "semanal",      0, 1),
            ("Lunes a Viernes",     "lun_vie",      0, 2),
            ("Quincenal (15 días)", "cada_15_dias", 1, 0),
            ("Mensual",             "mensual",      1, 1),
            ("Cada ? días",         "cada_x_dias",  1, 2),
        ]
        for txt, val, r, c in opciones:
            tk.Radiobutton(frame_rep, text=txt,
                           variable=self.var_repetir, value=val,
                           command=self._toggle_repeticion,
                           bg=self.GRIS_BG, font=("Segoe UI", 10)
                           ).grid(row=r, column=c, sticky="w", padx=10, pady=3)

        self.frame_x = tk.Frame(frame_rep, bg=self.GRIS_BG)
        self.frame_x.grid(row=2, column=0, columnspan=3, sticky="w", padx=10, pady=4)
        tk.Label(self.frame_x, text="¿Cada cuántos días? (1-999):",
                 bg=self.GRIS_BG, font=("Segoe UI", 10)).pack(side="left")
        vcmd = (self.register(lambda P: P.isdigit() and len(P) <= 3 or P == ""), "%P")
        self.ent_x = tk.Entry(self.frame_x, width=6, font=("Segoe UI", 11),
                              validate="key", validatecommand=vcmd)
        self.ent_x.insert(0, "7")
        self.ent_x.pack(side="left", padx=8)
        self.frame_x.grid_remove()

        # ══ 5. Acciones ═══════════════════════════════════════════
        self._seccion(f, "5. 🚀  Acciones")
        frame_btns = tk.Frame(f, bg=self.GRIS_BG)
        frame_btns.pack(padx=14, pady=6)

        self.btn_enviar = tk.Button(
            frame_btns, text="▶  ENVIAR AHORA",
            command=self._enviar,
            bg=self.VERDE_WA, fg="white",
            font=("Segoe UI", 11, "bold"),
            width=18, relief="flat", cursor="hand2", pady=6)
        self.btn_enviar.grid(row=0, column=0, padx=6)

        self.btn_programar = tk.Button(
            frame_btns, text="⏰  PROGRAMAR",
            command=self._programar,
            bg=self.VERDE_OSC, fg="white",
            font=("Segoe UI", 11, "bold"),
            width=16, relief="flat", cursor="hand2", pady=6)
        self.btn_programar.grid(row=0, column=1, padx=6)

        self.btn_detener = tk.Button(
            frame_btns, text="⏹  DETENER",
            command=self._detener,
            bg="#e74c3c", fg="white",
            font=("Segoe UI", 11, "bold"),
            width=12, relief="flat", cursor="hand2",
            pady=6, state="disabled")
        self.btn_detener.grid(row=0, column=2, padx=6)

        # ══ 6. Log ════════════════════════════════════════════════
        self._seccion(f, "6. 📋  Reporte de envíos")
        self.txt_log = scrolledtext.ScrolledText(
            f, height=9, width=80,
            font=("Consolas", 9),
            bg="#1e1e1e", fg="#d4d4d4",
            state="disabled")
        self.txt_log.pack(padx=14, pady=4)

        frame_rep_btns = tk.Frame(f, bg=self.GRIS_BG)
        frame_rep_btns.pack(anchor="w", padx=14, pady=4)
        tk.Button(frame_rep_btns, text="📂  Abrir reporte Excel",
                  command=self._abrir_reporte,
                  bg=self.VERDE_MED, fg="white",
                  font=("Segoe UI", 9), relief="flat",
                  cursor="hand2", padx=8, pady=4
                  ).pack(side="left", padx=4)
        tk.Button(frame_rep_btns, text="🗑  Limpiar log",
                  command=self._limpiar_log,
                  bg="#95a5a6", fg="white",
                  font=("Segoe UI", 9), relief="flat",
                  cursor="hand2", padx=8, pady=4
                  ).pack(side="left", padx=4)

        tk.Label(f, text="", bg=self.GRIS_BG).pack(pady=10)

    # ── Arrastre ──────────────────────────────────────────────────
    def _inicio_arrastre(self, event):
        self._drag_x = event.x_root - self.winfo_x()
        self._drag_y = event.y_root - self.winfo_y()

    def _arrastrar(self, event):
        self.geometry(f"+{event.x_root - self._drag_x}+{event.y_root - self._drag_y}")

    def _seccion(self, parent, titulo):
        tk.Label(parent, text=titulo,
                 font=("Segoe UI", 10, "bold"),
                 bg="#d5e8d4", fg="#0b5345",
                 anchor="w", padx=10, pady=5, relief="flat"
                 ).pack(fill="x", pady=(12, 2))

    # ── Lógica de UI ─────────────────────────────────────────────
    def _actualizar_grupos(self):
        for w in self.frame_grupos.winfo_children():
            w.destroy()
        self.grupos_vars.clear()
        grupos = grupos_disponibles(self.contactos)
        for i, g in enumerate(grupos):
            var = tk.BooleanVar(value=(g == "todos"))
            self.grupos_vars[g] = var
            r, c = divmod(i, 4)
            tk.Checkbutton(self.frame_grupos, text=g.capitalize(),
                           variable=var, bg=self.GRIS_BG,
                           font=("Segoe UI", 10)
                           ).grid(row=r, column=c, sticky="w", padx=12, pady=3)

    def _recargar_contactos(self):
        self.contactos = cargar_contactos()
        self._actualizar_grupos()
        self._log_ui(f"🔄 Contactos recargados: {len(self.contactos)} activos")

    def _toggle_datetime(self):
        estado = "disabled" if self.var_inmediato.get() else "normal"
        self.ent_fecha.config(state=estado)
        self.ent_hora.config(state=estado)

    def _toggle_repeticion(self):
        if self.var_repetir.get() == "cada_x_dias":
            self.frame_x.grid()
        else:
            self.frame_x.grid_remove()

    def _grupos_seleccionados(self) -> list:
        return [g for g, v in self.grupos_vars.items() if v.get()]

    def _filtrar_contactos(self, grupos: list) -> dict:
        if "todos" in grupos:
            return self.contactos
        return {t: d for t, d in self.contactos.items()
                if d.get("grupo", "").lower() in grupos}

    def _log_ui(self, msg: str):
        self.txt_log.config(state="normal")
        ts = datetime.now().strftime("%H:%M:%S")
        self.txt_log.insert("end", f"[{ts}] {msg}\n")
        self.txt_log.see("end")
        self.txt_log.config(state="disabled")

    def _limpiar_log(self):
        self.txt_log.config(state="normal")
        self.txt_log.delete("1.0", "end")
        self.txt_log.config(state="disabled")

    def _actualizar_contador(self):
        restantes = _envios_restantes()
        self.lbl_prueba.config(
            text=f"⚠️  VERSIÓN DE PRUEBA — {restantes} envíos restantes de {LIMITE_PRUEBA}")

    # ── Validaciones ─────────────────────────────────────────────
    def _validar(self) -> tuple:
        texto  = self.txt_mensaje.get("1.0", "end").strip()
        grupos = self._grupos_seleccionados()
        if not texto:
            messagebox.showwarning("AI SmartSend", "Escribe un mensaje antes de enviar.")
            return None, None
        if not grupos:
            messagebox.showwarning("AI SmartSend", "Selecciona al menos un grupo.")
            return None, None
        contactos_dest = self._filtrar_contactos(grupos)
        if not contactos_dest:
            messagebox.showwarning("AI SmartSend",
                "No hay contactos activos en los grupos seleccionados.\n"
                "Verifica tu archivo contactos.xlsx")
            return None, None
        restantes = _envios_restantes()
        if restantes <= 0:
            messagebox.showerror("AI SmartSend — Versión de Prueba",
                f"Has alcanzado el límite de {LIMITE_PRUEBA} envíos de prueba.\n\n"
                "Contacta al proveedor para adquirir la versión completa.")
            marcar_trial_agotado()
            return None, None
        if len(contactos_dest) > restantes:
            messagebox.showwarning("AI SmartSend — Versión de Prueba",
                f"Solo tienes {restantes} envío(s) restantes.\n"
                f"Se enviarán solo los primeros {restantes} contactos.")
            contactos_dest = dict(list(contactos_dest.items())[:restantes])
        elif restantes <= 3:
            messagebox.showinfo("AI SmartSend — Versión de Prueba",
                f"⚠️ Te quedan solo {restantes} envío(s) en la versión de prueba.")
        return texto, contactos_dest

    # ── Envío inmediato ───────────────────────────────────────────
    def _enviar(self):
        texto, contactos_dest = self._validar()
        if texto is None:
            return

        self.btn_enviar.config(state="disabled")
        self._log_ui(f"🚀 Iniciando envío a {len(contactos_dest)} contacto(s)...")
        self.update()

        def tarea():
            resultados = enviar_a_contactos(
                contactos_dest, texto, callback_log=self._log_ui)
            exitosos = sum(1 for r in resultados if r["estado"] == "Enviado")
            fallidos  = len(resultados) - exitosos
            _incrementar_contador(exitosos)
            self._log_ui(f"📊 Resultado: ✅ {exitosos} enviados | ❌ {fallidos} fallidos")
            guardar_reporte(resultados)
            self._log_ui("💾 Reporte guardado en reporte_envios.xlsx")
            self.btn_enviar.config(state="normal")
            self._actualizar_contador()

        threading.Thread(target=tarea, daemon=True).start()

    # ── Programar envío ───────────────────────────────────────────
    def _programar(self):
        texto, contactos_dest = self._validar()
        if texto is None:
            return

        if self.var_inmediato.get():
            messagebox.showinfo("AI SmartSend",
                "Para programar un envío futuro:\n\n"
                "1. Desmarca '⚡ Envío inmediato'\n"
                "2. Ingresa la fecha y hora deseada\n"
                "3. Vuelve a presionar PROGRAMAR")
            return

        try:
            dt_envio = datetime.strptime(
                f"{self.ent_fecha.get().strip()} {self.ent_hora.get().strip()}",
                "%d/%m/%Y %H:%M")
        except ValueError:
            messagebox.showerror("Error",
                "Formato incorrecto.\nUsa DD/MM/AAAA para fecha y HH:MM para hora.")
            return

        if dt_envio < datetime.now():
            messagebox.showwarning("AI SmartSend",
                "La fecha y hora indicadas ya pasaron.\nElige una fecha/hora futura.")
            return

        repetir  = self.var_repetir.get()
        hora_fmt = dt_envio.strftime("%H:%M")
        schedule.clear()

        def tarea_envio():
            self._log_ui("⏰ Ejecutando envío programado...")
            resultados = enviar_a_contactos(
                contactos_dest, texto, callback_log=self._log_ui)
            exitosos = sum(1 for r in resultados if r["estado"] == "Enviado")
            fallidos  = len(resultados) - exitosos
            _incrementar_contador(exitosos)
            self._log_ui(f"📊 Resultado: ✅ {exitosos} enviados | ❌ {fallidos} fallidos")
            guardar_reporte(resultados)
            self._log_ui("💾 Reporte guardado en reporte_envios.xlsx")
            self._actualizar_contador()
            self.btn_programar.config(state="normal")
            self.btn_detener.config(state="disabled")

        if repetir == "diario":
            schedule.every().day.at(hora_fmt).do(
                lambda: threading.Thread(target=tarea_envio, daemon=True).start())
            self._log_ui(f"⏰ Programado DIARIO a las {hora_fmt}")
        elif repetir == "cada_15_dias":
            schedule.every(15).days.at(hora_fmt).do(
                lambda: threading.Thread(target=tarea_envio, daemon=True).start())
            self._log_ui(f"⏰ Programado cada 15 días a las {hora_fmt}")
        elif repetir == "cada_x_dias":
            try:
                x = max(1, int(self.ent_x.get()))
            except ValueError:
                x = 7
            schedule.every(x).days.at(hora_fmt).do(
                lambda: threading.Thread(target=tarea_envio, daemon=True).start())
            self._log_ui(f"⏰ Programado cada {x} días a las {hora_fmt}")
        elif repetir == "semanal":
            dia_semana = ["monday", "tuesday", "wednesday", "thursday",
                          "friday", "saturday", "sunday"][dt_envio.weekday()]
            getattr(schedule.every(), dia_semana).at(hora_fmt).do(
                lambda: threading.Thread(target=tarea_envio, daemon=True).start())
            self._log_ui(f"⏰ Programado SEMANAL ({dia_semana}) a las {hora_fmt}")
        elif repetir == "lun_vie":
            for dia in ["monday", "tuesday", "wednesday", "thursday", "friday"]:
                getattr(schedule.every(), dia).at(hora_fmt).do(
                    lambda: threading.Thread(target=tarea_envio, daemon=True).start())
            self._log_ui("⏰ Programado LUNES A VIERNES a las {hora_fmt}")
        elif repetir == "mensual":
            dia_num = dt_envio.day
            schedule.every().day.at(hora_fmt).do(
                lambda: threading.Thread(target=tarea_envio, daemon=True).start()
                if datetime.now().day == dia_num else None)
            self._log_ui(f"⏰ Programado MENSUAL el día {dia_num} a las {hora_fmt}")

        self._iniciar_schedule()
        self.btn_programar.config(state="disabled")
        self.btn_detener.config(state="normal")

    def _iniciar_schedule(self):
        self._job_activo = True
        def loop():
            while self._job_activo:
                schedule.run_pending()
                time.sleep(20)
        self._hilo = threading.Thread(target=loop, daemon=True)
        self._hilo.start()

    def _detener(self):
        schedule.clear()
        self._job_activo = False
        self.btn_programar.config(state="normal")
        self.btn_detener.config(state="disabled")
        self._log_ui("⏹ Programación detenida.")

    def _abrir_reporte(self):
        import subprocess
        if EXCEL_REP.exists():
            subprocess.Popen(["start", "", str(EXCEL_REP)], shell=True)
        else:
            messagebox.showinfo("AI SmartSend",
                "Aún no hay reporte generado.\nRealiza un envío primero.")

    def _salir(self):
        if self._job_activo:
            if not messagebox.askyesno("AI SmartSend",
                "Hay envíos programados activos.\n¿Seguro que quieres salir?"):
                return
        self._job_activo = False
        schedule.clear()
        self.destroy()


# ══════════════════════════════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    app = App()
    app.mainloop()
